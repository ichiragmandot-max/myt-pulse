import openpyxl,collections,re,json,sys
SRC=sys.argv[1]
wb=openpyxl.load_workbook(SRC,data_only=True)
comb=wb['Teacher Combined Data']; scomb=wb['Staff Combined Data']
ts2src={}
for ws in (comb,scomb):
    for r in ws.iter_rows(min_row=2,values_only=True):
        if r[0]: ts2src.setdefault(r[0],set()).add((r[1],r[2]))
T=re.compile(r'^How clearly does (.+?) explain\s+(.+?)\??$')
S=re.compile(r'^How helpful was (.+?)\??$')
sub_fix={'Maths':'Math','Economic':'Economics'}
out={'overall':[], 'teacher':[], 'staff':[], 'unmapped':[]}
seen=set()
for ws in wb:
    if ws.title in ('Dashboard','Teacher Combined Data','Staff Combined Data','Form Links','Sheet1'): continue
    rows=list(ws.iter_rows(values_only=True))
    if len(rows)<2: continue
    h=[str(x).strip() if x is not None else '' for x in rows[0]]
    # identify centre/batch
    votes=collections.Counter()
    for r in rows[1:]:
        for cb in ts2src.get(r[0],()): votes[cb]+=1
    if not votes: out['unmapped'].append(ws.title); continue
    (centre,batch),_=votes.most_common(1)[0]
    if len(votes)>1: print('WARN multi',ws.title,votes,file=sys.stderr)
    blocks=[]
    i=0
    while i<len(h):
        m=T.match(h[i]); s=S.match(h[i])
        if m: blocks.append(('t',i,m.group(1).strip(),m.group(2).strip())); i+=6
        elif s: blocks.append(('s',i,s.group(1).strip(),None)); i+=6
        else: i+=1
    for r in rows[1:]:
        if not r[0]: continue
        key=(r[0],centre,batch)
        if key in seen: continue
        seen.add(key)
        rid=len(out['overall'])+1
        num=lambda v: float(v) if isinstance(v,(int,float)) else None
        txt=lambda v: str(v).strip() if v not in (None,'') else None
        out['overall'].append(dict(rid=rid,ts=str(r[0]),centre=centre,batch=batch,q1=num(r[1]),q2=num(r[2]),q3=num(r[3]),comment=txt(r[4])))
        for kind,c,name,subj in blocks:
            vals=[num(r[c+j]) if c+j<len(r) else None for j in range(5)]
            com=txt(r[c+5]) if c+5<len(r) else None
            if all(v is None for v in vals) and not com: continue
            rec=dict(rid=rid,centre=centre,batch=batch,name=name,r=vals,comment=com)
            if kind=='t': rec['subject']=sub_fix.get(subj,subj); out['teacher'].append(rec)
            else: out['staff'].append(rec)
print(len(out['overall']),len(out['teacher']),len(out['staff']),out['unmapped'],file=sys.stderr)
json.dump(out,open('raw2025.json','w'),default=str)
