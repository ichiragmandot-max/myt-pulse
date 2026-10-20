import json,collections,sys
from names import *
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table,TableStyleInfo
d=json.load(open('raw2025.json'))
NAVY='203864'; RED='DD152B'
wb=Workbook()
H=Font(bold=True,color='FFFFFF'); HF=PatternFill('solid',fgColor=NAVY)
def sheet(title,headers,widths):
    ws=wb.create_sheet(title); ws.append(headers)
    for i,w in enumerate(widths): ws.column_dimensions[chr(65+i)].width=w
    for c in ws[1]: c.font=H; c.fill=HF; c.alignment=Alignment(vertical='center',wrap_text=True)
    ws.row_dimensions[1].height=30; ws.freeze_panes='A2'; return ws
# README
rd=wb.active; rd.title='READ ME'
rd.column_dimensions['A'].width=110
lines=[('MYT Pulse — Feedback Roster 2026',Font(bold=True,size=16,color=NAVY)),
('Fill this in, then upload it on the admin page (Setup tab). You can re-upload any time before or during the week.',None),
('',None),
('Sheet "Batches"',Font(bold=True,color=RED)),
('  One row per centre + batch that should appear on the form. Put the number of active students so the dashboard can show response rate. Set Active = No to hide a batch.',None),
('Sheet "Teachers"',Font(bold=True,color=RED)),
('  One row per centre + batch + subject + teacher. A student sees only the teachers listed for their centre and batch.',None),
('  If a subject has two or more teachers in a batch (e.g. different divisions), the student is asked "Who teaches you Physics?" and picks.',None),
('  A teacher who takes two subjects in the same batch is rated once, with both subjects shown.',None),
('Sheet "Staff"',Font(bold=True,color=RED)),
('  One row per centre + batch + staff member. Write "All" in Batch to show that staff member to every batch at that centre.',None),
('Sheet "People"',Font(bold=True,color=RED)),
('  Every teacher and staff member, once. The Name must match exactly what you use in Teachers/Staff (use the dropdown).',None),
('  Phone: used only to prepare the SMS list with their PIN. PIN: leave blank and the system creates one.',None),
('  "2025 names": how this person was spelled in last year\'s forms, separated by commas. This is what links this year to last year.',None),
('  "Please check" column: questions from me about last year\'s names. Fix the 2025 names column if needed, then clear the note.',None),
('',None),
('Prefilled from the 2025 forms. Delete people who left and add new joiners.',Font(italic=True,color='666666')),
]
for t,f in lines:
    rd.append([t]);
    if f: rd.cell(rd.max_row,1).font=f
    rd.cell(rd.max_row,1).alignment=Alignment(wrap_text=True)
# Batches
cnt=collections.Counter((o['centre'],o['batch']) for o in d['overall'])
order=['Chembur Naka','Chembur Station','Ghatkopar','Matunga','Mahavir Nagar','Borivali','BARC']
bord=['6th ICSE','8th SSC','9th SSC','10th SSC','8th ICSE','9th ICSE','10th ICSE','8th CBSE','9th CBSE','11th Science','12th Science','11th Science GRIT','12th Science GRIT','Nishtha','11th Commerce','12th Commerce']
key=lambda cb:(order.index(cb[0]),bord.index(cb[1]) if cb[1] in bord else 99)
wsb=sheet('Batches',['Centre','Batch','Active students','Active (Yes/No)','2025 responses (info)'],[18,20,16,15,20])
for cb in sorted(cnt,key=key): wsb.append([cb[0],cb[1],None,'Yes',cnt[cb]])
# Teachers
ta=collections.Counter()
for t in d['teacher']:
    ta[(t['centre'],t['batch'],SUBJ.get(t['subject'],t['subject']),canon(t['name']))]+=1
wst=sheet('Teachers',['Centre','Batch','Subject','Teacher name','2025 ratings (info)'],[18,20,26,24,18])
for k in sorted(ta,key=lambda k:(key(k[:2]),k[2],k[3])): wst.append([*k,ta[k]])
# Staff
sa=collections.Counter((s['centre'],s['batch'],canon(s['name'])) for s in d['staff'])
wss=sheet('Staff',['Centre','Batch (or All)','Staff name','2025 ratings (info)'],[18,20,24,18])
# collapse to All if staff appear in every batch of that centre
batches_by_c=collections.defaultdict(set)
for c,b in cnt: batches_by_c[c].add(b)
bysc=collections.defaultdict(set)
for (c,b,n) in sa: bysc[(c,n)].add(b)
for (c,n) in sorted(bysc,key=lambda x:(order.index(x[0]),x[1])):
    bs=bysc[(c,n)]
    if bs>=batches_by_c[c]: wss.append([c,'All',n,sum(sa[(c,b,n)] for b in bs)])
    else:
        for b in sorted(bs,key=lambda b:bord.index(b) if b in bord else 99): wss.append([c,b,n,sa[(c,b,n)]])
# People
raw=collections.defaultdict(set); role=collections.defaultdict(set)
for t in d['teacher']: raw[canon(t['name'])].add(t['name']); role[canon(t['name'])].add('Teacher')
for s in d['staff']: raw[canon(s['name'])].add(s['name']); role[canon(s['name'])].add('Staff')
for n in list(raw): raw[n].add(n)
wsp=sheet('People',['Name','Type','Phone (for SMS)','PIN (blank = auto)','2025 names (comma separated)','Please check'],[24,16,16,16,40,90])
for n in sorted(raw,key=str.lower):
    r=role[n]; ty='Teacher & Staff' if len(r)==2 else r.pop()
    wsp.append([n,ty,None,None,', '.join(sorted(raw[n],key=lambda x:(x!=n,x))),CONFIRM.get(n)])
    if CONFIRM.get(n): wsp.cell(wsp.max_row,6).fill=PatternFill('solid',fgColor='FFF2CC')
for ws in (wsp,):
    for row in ws.iter_rows(min_row=2):
        row[5].alignment=Alignment(wrap_text=True,vertical='top')
# dropdowns
np_=wsp.max_row
for ws,col in ((wst,'D'),(wss,'C')):
    dv=DataValidation(type='list',formula1=f"=People!$A$2:$A${np_+200}",allow_blank=True,showErrorMessage=True,errorTitle='Unknown name',error='Pick a name from the People sheet (add new people there first).')
    ws.add_data_validation(dv); dv.add(f'{col}2:{col}3000')
for ws,col in ((wst,'A'),(wss,'A'),(wsb,'A')):
    dv=DataValidation(type='list',formula1='"'+','.join(order)+'"',allow_blank=True); ws.add_data_validation(dv); dv.add(f'{col}2:{col}3000')
dv=DataValidation(type='list',formula1='"Teacher,Staff,Teacher & Staff"'); wsp.add_data_validation(dv); dv.add('B2:B1000')
dv=DataValidation(type='list',formula1='"Yes,No"'); wsb.add_data_validation(dv); dv.add('D2:D500')
for ws in (wsb,wst,wss,wsp): ws.auto_filter.ref=ws.dimensions
wb.save(sys.argv[1])
print(wsb.max_row-1,'batches',wst.max_row-1,'teacher rows',wss.max_row-1,'staff rows',wsp.max_row-1,'people',sum(1 for n in raw if CONFIRM.get(n)),'flags')
